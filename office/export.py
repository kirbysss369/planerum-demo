"""Build editable office files from a validated program JSON revision."""

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor as WordRGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as SlideRGBColor
from pptx.util import Inches, Pt as SlidePt


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)


def fmt(hours):
    return f"{float(hours):g}" if isinstance(hours, (int, float)) else str(hours)


def set_document_metadata(properties, title, subject):
    timestamp = datetime.now(timezone.utc)
    properties.title = title
    properties.subject = subject
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.created = timestamp
    properties.modified = timestamp


def make_docx(program, output, coverage=None, regulatory=None):
    spec = program["spec"]
    doc = Document()
    set_document_metadata(doc.core_properties, spec["title"], "Образовательная программа")
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = WordRGBColor(0, 0, 0)
        ppr = style._element.pPr
        if ppr is not None:
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
    doc.add_paragraph("УТВЕРЖДЕНО: приказ № ______ от «___» __________ ____ г.")
    kind = (
        "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА ПОВЫШЕНИЯ КВАЛИФИКАЦИИ"
        if spec["kind"] == "qualification"
        else "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА"
    )
    doc.add_heading(kind, 0)
    doc.add_heading(spec["title"], 1)
    doc.add_paragraph(f"Аудитория: {spec['audience']}. Форма: {spec['delivery']}.")
    doc.add_paragraph(f"Объем: {fmt(spec['requested'])} академических часов.")
    doc.add_paragraph(f"Входной уровень: {spec['entry_level']}.")
    doc.add_paragraph(f"Итоговая аттестация: {spec['attestation']}.")
    doc.add_heading("1. Общая характеристика программы", 1)
    if regulatory:
        doc.add_heading("Нормативные основания подготовки", 2)
        for basis in regulatory.get("basis", []):
            doc.add_paragraph(f"{basis['title']} — {basis['scope']}. {basis['location']}")
        doc.add_paragraph(regulatory["note"])
    doc.add_heading("1.1. Цель и результаты обучения", 2)
    doc.add_paragraph(program["goal"])
    if coverage is None:
        add_table(doc, ["ID", "Результат обучения", "Темы", "Оценочные средства"], [
            [o["id"], o["text"], ", ".join(o.get("topic_ids", [])), ", ".join(o.get("assessment_ids", []))]
            for o in program["outcomes"]
        ])
    else:
        def labels(items):
            return "; ".join(f"{item['id']} — {item['text']}" for item in items) or "Не подготовлены"
        add_table(doc, ["Компетенция", "Темы", "Тестовые вопросы", "Практические задания"], [
            [f"{row['outcome_id']} — {row['outcome_text']}", labels(row["topics"]), labels(row["tests"]), labels(row["practice"])]
            for row in coverage
        ])
    doc.add_paragraph("Проверка связей подтверждает структуру. Корректность содержания и соответствие заданий компетенциям проверяет методист.")
    for outcome in program["outcomes"]:
        doc.add_heading(f"Карта компетенции {outcome['id']}", 3)
        doc.add_paragraph("Знания: " + "; ".join(outcome.get("knowledge", [])))
        doc.add_paragraph("Умения: " + "; ".join(outcome.get("skills", [])))
        doc.add_paragraph("Навыки: " + "; ".join(outcome.get("experience", [])))
        if outcome.get("standard_code"):
            source = next((s for s in program.get("sources", []) if s["id"] == outcome.get("source_id")), None)
            doc.add_paragraph("Код трудовой функции: " + outcome["standard_code"])
            if source:
                doc.add_paragraph(f"Основание: {source['title']} — {source['location']}")
                doc.add_paragraph("Подтверждающий фрагмент: " + source.get("evidence", ""))
        doc.add_paragraph("Темы: " + ", ".join(outcome.get("topic_ids", [])))
        doc.add_paragraph("Оценочные средства: " + ", ".join(outcome.get("assessment_ids", [])))
    landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = Inches(11.69), Inches(8.27)
    landscape.left_margin = landscape.right_margin = Inches(0.65)
    doc.add_heading("2. Учебный план", 1)
    module_rows = []
    topic_rows = []
    for module in program["modules"]:
        module_kinds = {k: 0 for k in ("lecture", "practice", "laboratory", "self_study")}
        for topic in module["topics"]:
            kinds = {k: 0 for k in ("lecture", "practice", "laboratory", "self_study")}
            for activity in topic["activities"]:
                kinds[activity["kind"]] += activity["hours"]
                module_kinds[activity["kind"]] += activity["hours"]
            topic_rows.append([module["title"], topic["title"], *[fmt(kinds[k]) for k in kinds], fmt(sum(kinds.values()))])
        module_rows.append([module["title"], *[fmt(module_kinds[k]) for k in module_kinds], fmt(sum(module_kinds.values()))])
    module_rows.append(["Итоговая аттестация", "0", "0", "0", "0", fmt(program["final_attestation_hours"])])
    topic_rows.append(["Итоговая аттестация", "", "0", "0", "0", "0", fmt(program["final_attestation_hours"])])
    add_table(doc, ["Модуль", "ЛК", "ПР", "ЛБ", "СРС", "Всего"], module_rows)
    doc.add_paragraph(f"Итого: {fmt(spec['requested'])} часов, включая итоговую аттестацию.")
    doc.add_heading("3. Учебно-тематический план", 1)
    add_table(doc, ["Модуль", "Тема", "ЛК", "ПР", "ЛБ", "СРС", "Всего"], topic_rows)
    doc.add_paragraph(f"Итого: {fmt(spec['requested'])} часов, включая итоговую аттестацию.")
    portrait = doc.add_section(WD_SECTION_START.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width, portrait.page_height = Inches(8.27), Inches(11.69)
    doc.add_heading("4. Рабочие программы и оценка качества", 1)
    lectures_by_topic = {lecture["topic_id"]: lecture for lecture in program.get("lectures", [])}
    for module in program["modules"]:
        doc.add_heading(module["title"], 2)
        for topic in module["topics"]:
            doc.add_heading(topic["title"], 3)
            doc.add_paragraph(topic.get("summary", ""))
            for activity in topic["activities"]:
                kind_names = {"lecture": "лекция", "practice": "практика", "laboratory": "лабораторная", "self_study": "самостоятельная работа"}
                kind = kind_names.get(activity["kind"], activity["kind"])
                doc.add_paragraph(f"{activity['name']} — {fmt(activity['hours'])} ч. ({kind}).")
            if topic["id"] in lectures_by_topic:
                doc.add_heading("План лекции", 3)
                doc.add_paragraph(lectures_by_topic[topic["id"]]["plan"])
    for assessment in program.get("assessments", []):
        doc.add_heading(f"Оценочное средство {assessment['id']}", 3)
        doc.add_paragraph(assessment["question"])
        if assessment.get("answer"):
            doc.add_paragraph("Ключ: " + assessment["answer"])
        if assessment.get("rubric"):
            doc.add_paragraph("Критерии: " + assessment["rubric"])
    doc.add_heading("5. Источники", 1)
    if program.get("sources"):
        for source in program["sources"]:
            status = "подтвержден" if source["status"] == "verified" else "требует проверки"
            doc.add_paragraph(f"{source['title']} — {source['location']} [{status}]")
    else:
        doc.add_paragraph("Подтвержденные источники не предоставлены. Заполнить после проверки методистом.")
    doc.add_paragraph("Преподаватель: ____________________    Подпись: ____________________")
    doc.save(output)


def make_pptx(program, output):
    presentation = Presentation()
    set_document_metadata(presentation.core_properties, program["spec"]["title"], "Лекционные материалы")
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    topic_titles = {topic["id"]: topic["title"] for module in program["modules"] for topic in module["topics"]}
    for lecture in program.get("lectures", []):
        slides = lecture.get("slides") or [lecture.get("plan", "")]
        for index, body in enumerate(slides, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            title = topic_titles.get(lecture["topic_id"], lecture["topic_id"])
            slide.shapes.title.text = f"{title} · {index}"
            for paragraph in slide.shapes.title.text_frame.paragraphs:
                paragraph.font.size = SlidePt(32)
                paragraph.font.color.rgb = SlideRGBColor(23, 31, 45)
            frame = slide.placeholders[1].text_frame
            frame.text = body
            for paragraph in frame.paragraphs:
                paragraph.font.size = SlidePt(24)
                paragraph.font.color.rgb = SlideRGBColor(23, 31, 45)
    if not presentation.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = program["spec"]["title"]
        slide.placeholders[1].text = "Лекционные занятия не предусмотрены"
    presentation.save(output)


def main():
    program = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    report = out / "validation-report.json"
    validation = json.loads(report.read_text(encoding="utf-8")) if report.exists() else {}
    make_docx(program, out / "program.docx", validation.get("coverage"), validation.get("regulatory"))
    make_pptx(program, out / "lectures.pptx")


if __name__ == "__main__":
    main()
