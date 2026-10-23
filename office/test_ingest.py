import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter

from ingest import validated_extract, UnsafeDocument


class ReferenceExtractionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def test_text_is_preserved(self):
        path = self.root / "reference.txt"
        text = "Демонстрационный референс: 72 часа, защита проекта."
        path.write_text(text, encoding="utf-8")
        self.assertEqual(validated_extract(path), text)

    def test_empty_and_oversized_text_are_rejected(self):
        path = self.root / "reference.txt"
        for text in (" \n\t", "x" * 200_001):
            path.write_text(text, encoding="utf-8")
            with self.assertRaises(ValueError):
                validated_extract(path)

    def test_docx_includes_table_contents(self):
        path = self.root / "reference.docx"
        document = Document()
        document.add_paragraph("Демонстрационный план")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Самостоятельная работа: 254 часа"
        document.save(path)
        text = validated_extract(path)
        self.assertIn("Демонстрационный план", text)
        self.assertIn("Самостоятельная работа: 254 часа", text)

    def test_archive_with_excessive_entries_is_rejected(self):
        path = self.root / "archive.docx"
        with ZipFile(path, "w") as archive:
            for index in range(2001):
                archive.writestr(str(index), "")
        with self.assertRaisesRegex(ValueError, "unpacked size limit"):
            validated_extract(path)

    def test_pdf_without_text_is_rejected(self):
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "no extractable text"):
            validated_extract(path)

    def test_pdf_page_limit_is_enforced(self):
        path = self.root / "long.pdf"
        writer = PdfWriter()
        for _ in range(101):
            writer.add_blank_page(width=100, height=100)
        writer.write(path)
        with self.assertRaisesRegex(ValueError, "100 pages"):
            validated_extract(path)

    def test_docx_rejects_macros_embedded_files_and_hidden_archives(self):
        for name, payload in (
            ("word/vbaProject.bin", b"macro"),
            ("word/embeddings/object.bin", b"embedded"),
            ("word/media/image.png", b"PK\x03\x04hidden archive"),
            ("word/hidden.xml", b'<!DOCTYPE x [<!ENTITY e "text">]><x>&e;</x>'),
        ):
            with self.subTest(name=name):
                path = self.root / "unsafe.docx"
                document = Document()
                document.add_paragraph("Обычный текст")
                document.save(path)
                with ZipFile(path, "a") as archive:
                    archive.writestr(name, payload)
                with self.assertRaises(UnsafeDocument):
                    validated_extract(path)

    def test_renamed_macro_enabled_docx_is_rejected(self):
        path = self.root / "renamed.docx"
        source = BytesIO()
        Document().save(source)
        with ZipFile(source) as original, ZipFile(path, "w") as archive:
            for entry in original.infolist():
                data = original.read(entry)
                if entry.filename == "[Content_Types].xml":
                    data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml", b"application/vnd.ms-word.document.macroEnabled.main+xml")
                archive.writestr(entry, data)
        with self.assertRaises(UnsafeDocument):
            validated_extract(path)

    def test_normal_hyperlinks_remain_allowed(self):
        path = self.root / "with-link.docx"
        document = Document()
        document.add_paragraph("Материал со ссылкой на источник")
        document.part.relate_to("https://example.org/reference", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        document.save(path)
        self.assertIn("Материал со ссылкой", validated_extract(path))

    def test_pdf_rejects_javascript_attachments_and_encryption(self):
        for kind in ("javascript", "attachment", "encrypted"):
            with self.subTest(kind=kind):
                path = self.root / (kind + ".pdf")
                writer = PdfWriter()
                writer.add_blank_page(width=100, height=100)
                if kind == "javascript":
                    writer.add_js("app.alert('test');")
                elif kind == "attachment":
                    writer.add_attachment("note.txt", b"attachment")
                else:
                    writer.encrypt("test-password")
                writer.write(path)
                with self.assertRaises(UnsafeDocument):
                    validated_extract(path)


if __name__ == "__main__":
    unittest.main()
