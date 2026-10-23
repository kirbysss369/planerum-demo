"""Extract bounded reference text; document contents are data, never commands."""

import json
import stat
import sys
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject


class UnsafeDocument(ValueError):
    pass


class ReferenceTooLarge(ValueError):
    pass


class NoReferenceText(ValueError):
    pass


def check_docx(path):
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
            raise ReferenceTooLarge("DOCX exceeds the unpacked size limit")
        seen = set()
        blocked_extensions = {".zip", ".7z", ".rar", ".gz", ".tar", ".xz", ".bz2", ".docx", ".xlsx", ".pptx", ".docm", ".exe", ".dll", ".js", ".vbs"}
        for item in entries:
            name = item.filename.casefold()
            relative = PurePosixPath(name)
            if name in seen or relative.is_absolute() or ".." in relative.parts or "\\" in name:
                raise UnsafeDocument("ambiguous archive path")
            seen.add(name)
            if item.flag_bits & 1 or stat.S_ISLNK(item.external_attr >> 16):
                raise UnsafeDocument("encrypted or linked entry")
            if item.is_dir():
                continue
            if name.startswith(("word/embeddings/", "word/activex/")) or "vbaproject" in name or relative.suffix in blocked_extensions:
                raise UnsafeDocument("active or embedded content")
            with archive.open(item) as member:
                signature = member.read(8)
            if signature.startswith((b"PK\x03\x04", b"PK\x05\x06", b"7z\xbc\xaf\x27\x1c", b"Rar!", b"\x1f\x8b")):
                raise UnsafeDocument("nested archive")
            if name.endswith(".xml") or name.endswith(".rels"):
                metadata = name == "[content_types].xml" or name.endswith(".rels")
                if metadata and item.file_size > 1 << 20:
                    raise ReferenceTooLarge("package metadata is too large")
                data = archive.read(item)
                if b"<!doctype" in data.lower() or b"<!entity" in data.lower():
                    raise UnsafeDocument("XML declarations are not allowed")
                if not metadata:
                    continue
                for element in ET.fromstring(data).iter():
                    content_type = element.attrib.get("ContentType", "").casefold()
                    relation = element.attrib.get("Type", "").rsplit("/", 1)[-1].casefold()
                    if any(part in content_type for part in ("macroenabled", "vbaproject", "oleobject", "activex")) or relation in {"oleobject", "package", "control", "vbaproject", "attachedtemplate"}:
                        raise UnsafeDocument("active package relationship")


def check_pdf(reader):
    if reader.is_encrypted:
        raise UnsafeDocument("encrypted PDF")
    pending, seen = [reader.trailer], set()
    examined = 0
    while pending:
        examined += 1
        if examined > 50000:
            raise ReferenceTooLarge("PDF structure is too complex")
        node = pending.pop()
        if isinstance(node, IndirectObject):
            key = (node.idnum, node.generation)
            if key in seen:
                continue
            seen.add(key)
            node = node.get_object()
        if isinstance(node, DictionaryObject):
            if any(key in node for key in ("/JavaScript", "/JS", "/EmbeddedFiles", "/EF", "/XFA", "/RichMediaContent")):
                raise UnsafeDocument("active or embedded PDF content")
            action = node.get("/S", "")
            if isinstance(action, IndirectObject):
                action = action.get_object()
            if str(action) in {"/JavaScript", "/Launch", "/SubmitForm", "/ImportData", "/Rendition", "/GoToR"}:
                raise UnsafeDocument("active PDF action")
            children = list(node.values())
        elif isinstance(node, ArrayObject):
            children = node
        else:
            continue
        if len(pending) + len(children) > 50000:
            raise ReferenceTooLarge("PDF structure is too complex")
        pending.extend(children)


def extract(path):
    suffix = path.suffix.lower()
    if suffix == ".docx":
        check_docx(path)
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise UnsafeDocument("encrypted PDF")
        if len(reader.pages) > 100:
            raise ReferenceTooLarge("PDF has more than 100 pages")
        check_pdf(reader)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    raise ValueError("unsupported reference format")


def validated_extract(path):
    text = extract(path)
    if not text.strip():
        raise NoReferenceText("reference contains no extractable text")
    if len(text) > 200_000:
        raise ReferenceTooLarge("extracted text is too large")
    return text


def main():
    # The importer runs in a separate Linux process with bounded address space.
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (512 << 20, 512 << 20))
    try:
        path = Path(sys.argv[1])
        if path.stat().st_size > 10 << 20:
            raise ReferenceTooLarge("reference file is too large")
        text = validated_extract(path)
        result = {"text": text, "characters": len(text)}
    except NoReferenceText:
        result = {"error_code": "no_text"}
    except UnsafeDocument:
        result = {"error_code": "unsafe_document"}
    except (ReferenceTooLarge, MemoryError):
        result = {"error_code": "reference_too_large"}
    except Exception:
        result = {"error_code": "reference_unreadable"}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
